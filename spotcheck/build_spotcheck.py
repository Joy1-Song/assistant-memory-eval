"""Build the blind human spot-check sheet for round 2 (judge-v2).

Random 20% of the 156 judged answers in outputs/live_r2 (8 per config, fixed seed), shown with exactly the
rubric the LLM judge saw. The judge's verdicts are NOT in the sheet; score_spotcheck.py joins them afterwards.
python spotcheck/build_spotcheck.py  ->  spotcheck/人工抽检表_第二轮.xlsx
"""

import csv
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from run_live import judge_values  # noqa: E402

SEED, PER_CONFIG = 20261025, 8
OUT = ROOT / "spotcheck/人工抽检表_第二轮.xlsx"


def sample_items():
    """The 32 sampled answers in sheet order, each with the rubric the judge saw (no judge verdict)."""
    bench = json.loads((ROOT / "data/memory_benchmark.json").read_text(encoding="utf-8"))
    probes = {p["id"]: (sc, p) for sc in bench["scenarios"] for p in sc["probes"]}
    rows = [r for r in csv.DictReader((ROOT / "outputs/live_r2/human_review.csv").open(encoding="utf-8-sig"))
            if r["judge_pass"] in {"True", "False"}]
    rng = random.Random(SEED)
    sample = []
    for cfg in ("none", "naive", "best", "best_global"):
        sample += rng.sample([r for r in rows if r["config"] == cfg], PER_CONFIG)
    rng.shuffle(sample)
    items = []
    for i, r in enumerate(sample, 1):
        sc, probe = probes[r["probe"]]
        need, avoid = judge_values(sc, probe)
        items.append({"n": i, "probe": r["probe"], "config": r["config"], "question": probe["question"],
                      "at": probe["at"].replace("T", " "), "no_memory": probe["no_memory"],
                      "memories": [m for m in r["memories_in_context"].split(" | ") if m],
                      "need": need, "avoid": avoid, "note": probe.get("judge_note") or "", "answer": r["answer"]})
    return items


def main():
    items = sample_items()

    wb = Workbook()
    guide = wb.active
    guide.title = "说明"
    for line in [
        "人工抽检：判断助理的回答是否通过（共 32 条，约 30–40 分钟）",
        "",
        "只在「抽检」页的「你的判定」一列下拉选择 通过 / 不通过，理由选填。其他列不要改。",
        "",
        "三条都满足才算「通过」：",
        "① 用上了「应该用上的信息」：意思用上即可，不要求原词。这一格写着「（无）」时，这条自动满足。",
        "② 没有依赖「不能用的信息」，并且符合「备注」。旧信息被更新后，新信息里仍保留的部分可以用（例如兴趣从科幻变成科幻和历史，推荐科幻书是对的）。",
        "③ 回答里出现的用户个人信息，必须来自「助理当时记得的内容」，并且和问题相关。凭空编造用户的情况，或者硬扯无关的个人信息，就不通过。一般性的安全提醒不算。",
        "",
        "只看和用户个人信息有关的这三点。和个人信息无关的常识错误（比如编了一个不存在的展览）不影响判定，可以写在理由里。",
        "判断依据是「助理当时记得的内容」：助理没被告知的信息，没用上不算它的错。",
        "这张表不显示模型评审的结论，做完我再对比，算一致率。",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True, size=13)

    ws = wb.create_sheet("抽检")
    head = ["序号", "问题", "提问时间", "这个问题和用户个人信息无关", "助理当时记得的内容", "应该用上的信息", "不能用的信息",
            "备注", "助理的回答", "你的判定", "理由（选填）", "probe", "config"]
    ws.append(head)
    for it in items:
        ws.append([it["n"], it["question"], it["at"], "是" if it["no_memory"] else "否",
                   "\n".join(it["memories"]) or "（无）", "\n".join(it["need"]) or "（无）",
                   "\n".join(it["avoid"]) or "（无）", it["note"], it["answer"], "", "", it["probe"], it["config"]])

    widths = [6, 22, 16, 10, 30, 22, 26, 18, 70, 12, 24, 8, 10]
    for col, w in zip("ABCDEFGHIJKLM", widths):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["L"].hidden = ws.column_dimensions["M"].hidden = True
    fill = PatternFill("solid", fgColor="FFF4C7")
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in range(2, len(items) + 2):
        ws[f"J{r}"].fill = fill
        ws[f"K{r}"].fill = fill
    dv = DataValidation(type="list", formula1='"通过,不通过"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{len(items) + 1}")
    ws.freeze_panes = "B2"
    wb.active = 1
    wb.save(OUT)
    print(f"{len(items)} rows -> {OUT}")


if __name__ == "__main__":
    main()
