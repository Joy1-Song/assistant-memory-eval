"""Compare the human verdicts with the LLM judge (judge-v2, round 2).

python spotcheck/score_spotcheck.py                    # from the filled Excel sheet
python spotcheck/score_spotcheck.py --from-json DIR    # from the web page's saved verdicts (DIR/verdicts/n01.json ...)
-> spotcheck/spotcheck_result.md
"""

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHEET = ROOT / "spotcheck/人工抽检表_第二轮.xlsx"
OUT = ROOT / "spotcheck/spotcheck_result.md"


def main():
    judge = {(r["probe"], r["config"]): r for r in
             csv.DictReader((ROOT / "outputs/live_r2/human_review.csv").open(encoding="utf-8-sig"))}
    if "--from-json" in sys.argv:
        folder = Path(sys.argv[sys.argv.index("--from-json") + 1]) / "verdicts"
        rows = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(folder.glob("*.json"))]
        rows = [{"n": r["n"], "probe": r["probe"], "config": r["config"], "verdict": r.get("verdict", ""),
                 "note": r.get("note", "")} for r in (x.get("data", x) for x in rows)]
        from build_spotcheck import sample_items
        questions = {it["n"]: it["question"] for it in sample_items()}
        seen = {r["n"] for r in rows}
        missing = [n for n in questions if n not in seen]
    else:
        ws = load_workbook(SHEET)["抽检"]
        head = [c.value for c in ws[1]]
        col = {name: head.index(name) for name in ("序号", "问题", "你的判定", "理由（选填）", "probe", "config")}
        rows, questions, missing = [], {}, []
        for row in ws.iter_rows(min_row=2, values_only=True):
            questions[row[col["序号"]]] = row[col["问题"]]
            rows.append({"n": row[col["序号"]], "probe": row[col["probe"]], "config": row[col["config"]],
                         "verdict": (row[col["你的判定"]] or "").strip(), "note": row[col["理由（选填）"]] or ""})
    done = []
    for r in rows:
        if r["verdict"] not in {"通过", "不通过"}:
            missing.append(r["n"])
            continue
        j = judge[(r["probe"], r["config"])]
        done.append({"n": r["n"], "probe": r["probe"], "config": r["config"], "q": questions[r["n"]],
                     "human": r["verdict"] == "通过", "judge": j["judge_pass"] == "True",
                     "judge_reason": j["judge_reason"], "note": r["note"]})
    missing = sorted(set(missing))

    n = len(done)
    agree = sum(d["human"] == d["judge"] for d in done)
    tp = sum(d["human"] and d["judge"] for d in done)
    tn = sum(not d["human"] and not d["judge"] for d in done)
    judge_only = sum(d["judge"] and not d["human"] for d in done)
    human_only = sum(d["human"] and not d["judge"] for d in done)
    p_o = agree / n if n else 0
    p_e = (((tp + judge_only) * (tp + human_only)) + ((tn + human_only) * (tn + judge_only))) / n ** 2 if n else 0
    kappa = (p_o - p_e) / (1 - p_e) if n and p_e < 1 else float("nan")

    lines = ["# 人工抽检：judge-v2 与人工判定的一致率（第二轮）", "",
             f"抽样：outputs/live_r2 的 156 条已评审回答中，按配置各随机 8 条（seed 20261025），共 32 条；"
             f"已判定 {n} 条，未判定 {len(missing)} 条 {missing if missing else ''}", "",
             f"- 一致：**{agree}/{n}**（{p_o:.0%}），Cohen's kappa {kappa:.2f}",
             f"- 两者都判通过 {tp}，都判不通过 {tn}；评审判通过、人判不通过 {judge_only}；人判通过、评审判不通过 {human_only}", "",
             "## 不一致的条目", "", "| 序号 | 用例 | 配置 | 问题 | 人工 | 评审 | 评审理由 | 人工理由 |",
             "|---:|---|---|---|---|---|---|---|"]
    for d in done:
        if d["human"] != d["judge"]:
            lines.append(f"| {d['n']} | {d['probe']} | {d['config']} | {d['q']} | {'通过' if d['human'] else '不通过'} | "
                         f"{'通过' if d['judge'] else '不通过'} | {d['judge_reason']} | {d['note']} |")
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()
